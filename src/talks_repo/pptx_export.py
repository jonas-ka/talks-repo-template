"""PPTX export for Google Slides, from the compiled (UA-1) PDF of a deck.

Google Slides imports neither Typst, PDF nor SVG, but it does import PPTX (File > Import
slides). Two modes:

* **hybrid** (default): every page is rebuilt with live text. Each text line of the PDF
  becomes a native text box at the same position (font, size, weight, colour kept; STIX
  Two Text is on Google Fonts), each photo or raster figure a picture shape at its
  rectangle, and only what is left (rules, panels, diagrams, vector plots, logos) is
  rendered as the slide background at `ppi`. Text is searchable, selectable and reusable;
  the vector artwork is a picture, which Slides could not represent faithfully anyway.
* **flat** (`--flat`): every page as one full-slide picture. For posters and pages the
  hybrid mode cannot take apart.

    talks pptx <slug>                       talks/<slug>/YYMM.<Event>-YourName.pptx
    talks pptx <slug> --pages 1,3-4         a subset of pages
    talks pptx <slug> --flat                pictures only
    talks pptx <slug> --png                 also keep the PNGs (talks/<slug>/png/)

How the hybrid mode reads the PDF (PyMuPDF):

* Typst embeds every font as a subset without `OS/2` and `cmap` tables, and a variable
  font (STIX Two Text on macOS) gives regular and bold the *same* name. Weight is
  recovered from the glyph advances: of two same-named subsets, the one whose shared
  characters are wider is the bold one (`_classify_fonts`). Each subset is renamed in
  memory (`Family-Bold`, `Family-Italic`, ...) so the text extraction can tell them apart.
* Text boxes are 10 % wider than their line, extended to the right, with the text
  left-aligned (the author, 2026-10-07): Google Slides ignores "do not wrap" and lays text out with
  slightly different metrics, so a box exactly as wide as the PDF line pushed its last word
  onto a second line after import. The extra width is empty space to the right; the text
  starts where it did. (`BOX_EXTRA`)
* Text boxes get zero insets and no wrap; their top is the baseline minus the font's
  ascent (hhea = typo metrics for STIX Two Text). Checked against PowerPoint for Mac on
  2026-09-29: 61 of 62 lines within 1.5 pt horizontally, baselines a uniform 1.8 pt lower
  there (`baseline` nudges them).
* Images are placed only when drawn unrotated and uncropped (their rectangle has the
  image's aspect ratio); anything else stays in the background.
* The background is the page with those text lines and images redacted away
  (`apply_redactions`, keeping line art), rendered at `ppi`.

Limits: one text box per line (no flowing paragraphs), equations arrive as STIX Two Math
glyph text, and the Typst source in the repo remains the editable original.
"""

from __future__ import annotations

import datetime as dt
import io
import re
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

import pymupdf
from fontTools.ttLib import TTFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu, Pt

from talks_repo import REPO_ROOT

EMU_PER_PT = 12700
EMU_PER_INCH = 914400

# PDF font family -> name Google Slides / PowerPoint know. Anything else is split at the
# CamelCase joints ("HelveticaNeue" -> "Helvetica Neue").
FONT_NAMES = {
    "STIXTwoText": "STIX Two Text",
    "STIXTwoMath": "STIX Two Math",
    "DejaVuSansMono": "Roboto Mono",      # Typst's raw font; Roboto Mono has the same 0.6 em advance
    "NewCMMath": "STIX Two Math",
    "LibertinusSerif": "Libertinus Serif",
}
BOLD_RATIO = 1.04    # mean advance of a bold subset over the regular one (measured: 1.09)


# ---------------------------------------------------------------- helpers
def parse_pages(spec: str | None, n: int) -> list[int]:
    """'1,3-4' -> [1, 3, 4] (1-based, clipped to the document); None -> all pages."""
    if not spec:
        return list(range(1, n + 1))
    pages: list[int] = []
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            a, b = part.split("-", 1)
            pages.extend(range(int(a), int(b) + 1))
        else:
            pages.append(int(part))
    return [p for p in pages if 1 <= p <= n]


def deck_pdf(slug_or_path: str) -> Path:
    """The deck's PDF: a path to a .pdf, or the newest *.pdf in talks/<slug>/."""
    p = Path(slug_or_path)
    if p.suffix.lower() == ".pdf":
        return p.resolve()
    folder = REPO_ROOT / "talks" / slug_or_path
    pdfs = sorted(folder.glob("*.pdf"), key=lambda q: q.stat().st_mtime, reverse=True)
    if not pdfs:
        raise FileNotFoundError(f"no PDF in {folder}; compile the deck first")
    return pdfs[0]


def display_font(family: str) -> str:
    if family in FONT_NAMES:
        return FONT_NAMES[family]
    if family.startswith("."):      # a macOS system fallback (".SFNS": glyphs like ✓ ✗): any sans will do
        return "Arial"
    return re.sub(r"(?<=[a-z])(?=[A-Z])|(?<=[A-Z])(?=[A-Z][a-z])", " ", family)


def emu(pt: float) -> Emu:
    return Emu(int(round(pt * EMU_PER_PT)))


# ---------------------------------------------------------------- fonts
@dataclass
class FontInfo:
    xref: int
    family: str            # PDF family without subset prefix and style suffix
    bold: bool
    italic: bool
    ascent: float          # em, from hhea
    descent: float         # em, negative
    advances: list[float] = field(default_factory=list)   # em per glyph id
    tag: str = ""          # the name the text extraction will report


def _style_from_name(base: str) -> tuple[str, bool, bool]:
    name = base.split("+")[-1]
    family, _, style = name.partition("-")
    style_l = style.lower()
    bold = "bold" in style_l or "black" in style_l or "heavy" in style_l
    italic = "italic" in style_l or "oblique" in style_l
    if not style and family.lower().endswith(("bold", "italic")):
        pass
    return family, bold, italic


def _rename_font(doc: pymupdf.Document, xref: int, name: str) -> None:
    """Set /BaseFont on a font and its descendants (in memory only; the file is not saved)."""
    doc.xref_set_key(xref, "BaseFont", "/" + name)
    kind, val = doc.xref_get_key(xref, "DescendantFonts")
    if kind == "array":
        for m in re.finditer(r"(\d+) 0 R", val):
            doc.xref_set_key(int(m.group(1)), "BaseFont", "/" + name)


def _classify_fonts(doc: pymupdf.Document) -> dict[str, FontInfo]:
    """Read every embedded font, recover weight from glyph advances, rename in memory.

    Returns the fonts by the name `get_text` will report for their spans."""
    fonts: dict[int, FontInfo] = {}
    for page in doc:
        for f in page.get_fonts(full=True):
            xref, base = f[0], f[3]
            if xref in fonts:
                continue
            family, bold, italic = _style_from_name(base)
            ascent, descent, adv = 0.8, -0.2, []
            try:
                tt = TTFont(io.BytesIO(doc.extract_font(xref)[3]))
                upem = tt["head"].unitsPerEm
                ascent, descent = tt["hhea"].ascent / upem, tt["hhea"].descent / upem
                adv = [tt["hmtx"].metrics[g][0] / upem for g in tt.getGlyphOrder()]
            except Exception:
                pass
            fi = FontInfo(xref, family, bold, italic, ascent, descent, adv, tag=f"{family}~{xref}")
            _rename_font(doc, xref, fi.tag)
            fonts[xref] = fi
    by_tag = {fi.tag: fi for fi in fonts.values()}

    # Advance per unicode character for every subset, from the glyph ids in the text trace.
    adv: dict[str, dict[int, float]] = defaultdict(dict)
    for page in doc:
        for span in page.get_texttrace():
            fi = by_tag.get(span["font"])
            if fi is None:
                continue
            for ch in span["chars"]:
                uni, gid = ch[0], ch[1]
                if gid < len(fi.advances):
                    adv[fi.tag][uni] = fi.advances[gid]
    # Same family and style, several subsets: the wider one is bold.
    groups: dict[tuple[str, bool, bool], list[FontInfo]] = defaultdict(list)
    for fi in fonts.values():
        groups[(fi.family, fi.bold, fi.italic)].append(fi)
    for members in groups.values():
        if len(members) < 2:
            continue
        def mean_adv(fi: FontInfo, chars: set[int]) -> float:
            return sum(adv[fi.tag][u] for u in chars) / max(len(chars), 1)
        base = min(members, key=lambda fi: sum(adv[fi.tag].values()) / max(len(adv[fi.tag]), 1))
        for fi in members:
            if fi is base:
                continue
            shared = set(adv[fi.tag]) & set(adv[base.tag])
            if shared and mean_adv(fi, shared) / mean_adv(base, shared) > BOLD_RATIO:
                fi.bold = True
    # Final names: Family, Family-Bold, Family-Italic, Family-BoldItalic (unique per xref
    # through a suffix the reader strips again).
    result: dict[str, FontInfo] = {}
    for fi in fonts.values():
        style = ("Bold" if fi.bold else "") + ("Italic" if fi.italic else "")
        fi.tag = fi.family + ("-" + style if style else "") + f"~{fi.xref}"
        _rename_font(doc, fi.xref, fi.tag)
        result[fi.tag] = fi
    return result


# ---------------------------------------------------------------- page parts
@dataclass
class Run:
    text: str
    font: FontInfo
    size: float
    color: tuple[int, int, int]
    superscript: bool
    link: str | None = None


@dataclass
class TextBox:
    x0: float
    x1: float
    baseline: float
    runs: list[Run]

    @property
    def ascent(self) -> float:
        return max(r.font.ascent * r.size for r in self.runs)

    @property
    def descent(self) -> float:
        return max(-r.font.descent * r.size for r in self.runs)

    @property
    def rect(self) -> pymupdf.Rect:
        return pymupdf.Rect(self.x0, self.baseline - self.ascent, self.x1, self.baseline + self.descent)


def _color(c: int) -> tuple[int, int, int]:
    return ((c >> 16) & 255, (c >> 8) & 255, c & 255)


def extract_text(page: pymupdf.Page, fonts: dict[str, FontInfo]) -> tuple[list[TextBox], list[pymupdf.Rect]]:
    """Text boxes for the slide (one per line, split at wide gaps and baseline jumps) and
    the rectangles to redact from the background. Rotated text stays in the background."""
    links = [(pymupdf.Rect(lk["from"]), lk.get("uri")) for lk in page.get_links() if lk.get("uri")]
    boxes: list[TextBox] = []
    for block in page.get_text("dict", flags=pymupdf.TEXT_PRESERVE_WHITESPACE)["blocks"]:
        if block["type"] != 0:
            continue
        for line in block["lines"]:
            if abs(line["dir"][0] - 1) > 1e-3 or abs(line["dir"][1]) > 1e-3:
                continue   # rotated: not a text box
            current: TextBox | None = None
            for span in line["spans"]:
                text = span["text"]
                if not text.strip():
                    # A space-only span (Typst emits them between differently styled words):
                    # keep it in the previous run's font, never as a box of its own.
                    if current is not None and abs(span["origin"][1] - current.baseline) <= 0.5:
                        current.runs[-1].text += text
                        current.x1 = max(current.x1, span["bbox"][2])
                    continue
                fi = fonts.get(span["font"])
                if fi is None:   # a font we could not read: fall back to the name
                    family, bold, italic = _style_from_name(span["font"])
                    fi = FontInfo(0, family, bold or bool(span["flags"] & 16), italic or bool(span["flags"] & 2), 0.8, -0.2)
                    fonts[span["font"]] = fi
                x0, y0, x1, y1 = span["bbox"]
                baseline = span["origin"][1]
                run = Run(text, fi, span["size"], _color(span["color"]), bool(span["flags"] & 1))
                for r, uri in links:
                    if r.intersects(pymupdf.Rect(span["bbox"])):
                        run.link = uri
                gap = x0 - current.x1 if current else 0
                if current is None or abs(baseline - current.baseline) > 0.5 or gap > 0.6 * span["size"]:
                    current = TextBox(x0, x1, baseline, [run])
                    boxes.append(current)
                else:
                    # Typst spaces words by positioning, so a word gap between two spans has
                    # no space character; put one back so the words do not run together.
                    prev = current.runs[-1]
                    if gap > 0.12 * span["size"] and not prev.text.endswith(" ") and not text.startswith(" "):
                        prev.text += " "
                    current.runs.append(run)
                    current.x1 = max(current.x1, x1)
    # Plain PDF text line rectangles (glyph boxes) for the redaction.
    redact = [b.rect for b in boxes]
    return boxes, redact


@dataclass
class Picture:
    rect: pymupdf.Rect
    data: bytes
    ext: str


def extract_images(doc: pymupdf.Document, page: pymupdf.Page) -> list[Picture]:
    """Images drawn unrotated and uncropped, with their native bytes."""
    pics: list[Picture] = []
    for info in page.get_image_info(xrefs=True):
        a, b, c, d_, e, f = info["transform"]
        if abs(b) > 1e-6 or abs(c) > 1e-6 or a <= 0 or d_ <= 0:
            continue   # rotated or flipped: leave it in the background
        w, h = info["width"], info["height"]
        rect = pymupdf.Rect(info["bbox"])
        if w == 0 or h == 0 or rect.is_empty:
            continue
        if abs((rect.width / rect.height) / (w / h) - 1) > 0.03:
            continue   # cropped (fit: cover) or stretched: keep in the background
        xref = info["xref"]
        try:
            raw = doc.extract_image(xref)
        except Exception:
            continue
        if raw.get("smask") or raw.get("ext") not in ("png", "jpeg", "jpg"):
            pix = pymupdf.Pixmap(doc, xref)
            if raw.get("smask"):
                pix = pymupdf.Pixmap(pix, pymupdf.Pixmap(doc, raw["smask"]))
            if pix.colorspace and pix.colorspace.n > 3:
                pix = pymupdf.Pixmap(pymupdf.csRGB, pix)
            data, ext = pix.tobytes("png"), "png"
        else:
            data, ext = raw["image"], raw["ext"]
        pics.append(Picture(rect, data, ext))
    return pics


def background_png(doc: pymupdf.Document, pno: int, text_rects: list[pymupdf.Rect],
                   image_rects: list[pymupdf.Rect], ppi: int) -> bytes:
    """The page without its text and placed images, rendered at `ppi`."""
    tmp = pymupdf.open()
    tmp.insert_pdf(doc, from_page=pno, to_page=pno)
    pg = tmp[0]
    if text_rects:
        for r in text_rects:
            pg.add_redact_annot(r + (-0.5, 0, 0.5, 0), fill=False)
        pg.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE, graphics=pymupdf.PDF_REDACT_LINE_ART_NONE,
                            text=pymupdf.PDF_REDACT_TEXT_REMOVE)
    if image_rects:
        for r in image_rects:
            pg.add_redact_annot(r, fill=False)
        pg.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_REMOVE, graphics=pymupdf.PDF_REDACT_LINE_ART_NONE,
                            text=pymupdf.PDF_REDACT_TEXT_NONE)
    pix = pg.get_pixmap(dpi=ppi, alpha=False)
    data = pix.tobytes("png")
    tmp.close()
    return data


# ---------------------------------------------------------------- pptx
BOX_EXTRA = 0.10   # text boxes this much wider than their line, to the right (Google Slides re-wraps otherwise)


def _add_textbox(slide, box: TextBox, baseline_shift: float = 0.0) -> None:
    top = box.baseline - box.ascent + baseline_shift
    height = box.ascent + box.descent
    width = (box.x1 - box.x0 + 2) * (1 + BOX_EXTRA)
    shape = slide.shapes.add_textbox(emu(box.x0), emu(top), emu(width), emu(height))
    tf = shape.text_frame
    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT        # explicit: the extra width stays on the right
    p.line_spacing = 1.0
    for run in box.runs:
        r = p.add_run()
        r.text = run.text
        r.font.name = display_font(run.font.family)
        r.font.size = Pt(run.size)
        r.font.bold = run.font.bold
        r.font.italic = run.font.italic
        r.font.color.rgb = RGBColor(*run.color)
        if run.superscript:
            r.font._element.set("baseline", "30000")
        if run.link:
            r.hyperlink.address = run.link
    shape.name = "text: " + box.runs[0].text[:30]


def export_pptx(pdf: Path, out: Path | None = None, ppi: int = 300, pages: str | None = None,
                keep_png: bool = False, flat: bool = False, baseline: float = 0.0,
                source: str | None = None) -> dict:
    """Write the PPTX; hybrid (live text + pictures + background) unless `flat`.

    `baseline` moves every text box up (negative) or down in points. The default puts the
    first baseline where the font's typo metrics say (browsers, so Google Slides); PowerPoint
    for Mac renders it 1.8 pt lower (measured 2026-09-29), so `-1.8` matches that."""
    doc = pymupdf.open(pdf)
    wanted = parse_pages(pages, len(doc))
    if not wanted:
        raise ValueError(f"no pages selected from {pdf.name} ({len(doc)} pages)")
    out = out or pdf.with_suffix(".pptx")
    png_dir = pdf.parent / "png"
    if keep_png:
        png_dir.mkdir(exist_ok=True)
    fonts = {} if flat else _classify_fonts(doc)

    first = doc[wanted[0] - 1].rect
    prs = Presentation()
    prs.slide_width = emu(first.width)
    prs.slide_height = emu(first.height)
    blank = prs.slide_layouts[6]

    src = source or (str(pdf.relative_to(REPO_ROOT)) if pdf.is_relative_to(REPO_ROOT) else str(pdf))
    stamp = dt.date.today().isoformat()
    written: list[Path] = []
    n_text = n_pics = 0
    px = (0, 0)
    for n in wanted:
        page = doc[n - 1]
        slide = prs.slides.add_slide(blank)
        if flat:
            boxes, pics, redact = [], [], []
            pix = page.get_pixmap(dpi=ppi, alpha=False)
            bg = pix.tobytes("png")
            px = (pix.width, pix.height)
        else:
            boxes, redact = extract_text(page, fonts)
            pics = extract_images(doc, page)
            bg = background_png(doc, n - 1, redact, [p.rect for p in pics], ppi)
        # Background first, so it sits behind everything.
        bg_shape = slide.shapes.add_picture(io.BytesIO(bg), 0, 0, width=prs.slide_width, height=prs.slide_height)
        bg_shape.name = f"background page {n}"
        if keep_png:
            png = png_dir / f"{pdf.stem}-p{n:02d}{'' if flat else '-background'}.png"
            png.write_bytes(bg)
            written.append(png)
        for pic in pics:
            shp = slide.shapes.add_picture(io.BytesIO(pic.data), emu(pic.rect.x0), emu(pic.rect.y0),
                                           width=emu(pic.rect.width), height=emu(pic.rect.height))
            shp.name = f"image {pic.rect.width:.0f}x{pic.rect.height:.0f} pt"
        for box in boxes:
            _add_textbox(slide, box, baseline)
        n_text += len(boxes)
        n_pics += len(pics)
        mode = "picture of the page" if flat else f"{len(boxes)} text boxes, {len(pics)} pictures, vector art as background"
        slide.notes_slide.notes_text_frame.text = (
            f"Source: {src}, page {n} of {len(doc)} ({mode}). Exported {stamp} by `talks pptx`; "
            f"the Typst source is the original.")
    prs.save(out)
    return {"pptx": out, "slides": len(wanted), "pages": len(doc), "px": px, "text_boxes": n_text,
            "pictures": n_pics, "fonts": sorted({display_font(f.family) for f in fonts.values()}),
            "bytes": out.stat().st_size, "png": written, "flat": flat}
