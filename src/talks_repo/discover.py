"""Stage C: file discovery, deck detection, pairing, and matching to talks.

Walks the read-only input folders, records every deck-like file in files.yaml,
scores each file for "is this a slide deck", groups the same deck in different
forms (Keynote/PPTX/Slides source + PDF export), and matches groups to
talks.yaml entries by date and text similarity. Never writes into the inputs.
"""

import json
import os
import re
import zipfile
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

from rapidfuzz import fuzz

from talks_repo import INPUT_ROOTS, InputRoot

DECK_EXT = {".key", ".pptx", ".ppt", ".pdf", ".gslides"}
SOURCE_EXT = (".key", ".pptx", ".gslides", ".ppt")  # preference order for source_file

_POSITIVE_WORDS = ("talk", "slides", "seminar", "colloquium", "presentation", "lecture", "keynote")
_OWN_RE = re.compile(r"karthein|jonas|(?<![a-z])jk(?![a-z])")
# Subfolders that hold other people's decks (templates, downloads, saved talks).
_FOREIGN_DIRS = ("vorlagen", "template", "other", "interesting", "save", "download", "received",
                 "example", "reference", "literature")
# Subfolders that hold the author's own deck for the trip.
_TALK_DIRS = ("talk", "presentation", "vortrag", "pitch", "slides", "seminar", "colloquium")
_NEGATIVE_WORDS = (
    "invoice", "itinerary", "boarding", "hotel", "receipt", "abstract", "poster", "flyer",
    "visa", "ticket", "reimburs", "agenda", "program", "schedule", "letter", "certificate",
    "registration", "booking", "confirmation", "map", "menu", "form", "cv", "resume",
    "proposal", "report", "paper", "manuscript", "arxiv", "handout", "timetable", "badge",
    "travel", "expense", "transactions", "training", "duty", "thesis", "bachelorthesis",
    "baggage", "luggage", "airbnb", "uber", "taxi", "photobook", "route", "payment", "flight",
    "rechnung", "quittung", "beleg", "bahn", "flug", "reise", "antrag",
)
# Folder-name acronyms expanded so they can match the CV's spelled-out event names.
_ACRONYMS = {
    "doe": "department of energy", "aps": "american physical society",
    "dpg": "german physical society spring conference", "smuk": "german physical society",
    "lecm": "low energy community meeting", "nnpss": "national nuclear physics summer school",
    "grc": "gordon research conference", "tcp": "trapped charged particles",
    "mats": "measurements with an advanced trapping system", "lbnl": "lawrence berkeley national laboratory",
    "anl": "argonne", "ornl": "oak ridge national laboratory", "wwnd": "winter workshop nuclear dynamics",
    "mipep": "mitchell institute physics enhancement program", "iqse": "iqse seminar", "isqe": "iqse seminar",
    "cycl": "cyclotron", "weh": "wilhelm und else heraeus dissertation award", "int": "int parity nonconservation",
    "genco": "nustar", "gain": "gain", "lns": "mit laboratory nuclear science", "hpg": "journal club",
    "kollaborationstreffen": "collaboration meeting", "arbeitstreffen": "arbeitstreffen kernphysik",
    "eurorib": "eurorib", "ak": "arbeitstreffen kernphysik", "funsymm": "long range plan fundamental symmetries",
    "pretownhall": "long range plan", "interview": "seminar colloquium", "fs": "fundamental symmetries",
    "radis": "radis workshop", "namo": "namo workshop", "ariel": "ariel science workshop",
    "eunpc": "eunpc22", "platan": "platan24", "dnp": "dnp conference american physical society",
    "woodnext": "woodnext site visit", "artsci": "artsci quantum research meeting",
    "humboldt": "humboldt networking conference", "becola": "rise becola collaboration meeting",
    "gentner": "gentner day", "verteidigung": "defense", "dissertation": "dissertation award symposium",
}
_TEMPLATE_WORDS = ("template", "master", "blank", "logo")
_VERSION_TOKENS = re.compile(
    r"\b(final|v\d+|copy|backup|reduced|smaller|filesize|gslide|short|long|old|new|draft|"
    r"prelim|preliminary|handout|notes|compressed|small|large|export)\b"
)
_DATE_PREFIX = re.compile(r"^(\d{2})(\d{2})(?:[.\-_](\d{1,2}))?[.\-_ ]")
_NAME_TOKENS = re.compile(r"\b(jonas|karthein|jonaskarthein|jk)\b")


@dataclass
class FileRecord:
    root: str
    path: str  # relative to the root
    folder: str  # first path component under the root, "" for loose files
    type: str
    mtime: str
    size: int
    cloud_only: bool
    analyzed: bool = False
    pages: int | None = None
    first_title: str | None = None
    landscape: bool | None = None
    aspect: float | None = None
    chars_per_page: float | None = None
    page_size: str | None = None  # "WxH" in points, first page
    gdrive_id: str | None = None
    own: bool = False  # the author's own deck (vs. someone else's saved in the trip folder)
    is_deck: float = 0.0
    deck_class: str = "unknown"
    reasons: list[str] = field(default_factory=list)
    group_id: str | None = None
    talk_id: str | None = None
    match_confidence: float | None = None
    status: str = "unreviewed"
    flags: list[str] = field(default_factory=list)

    @property
    def ref(self) -> str:
        return f"{self.root}:{self.path}"

    def to_entry(self) -> dict[str, Any]:
        return {
            "path": self.path,
            "root": self.root,
            "folder": self.folder,
            "type": self.type,
            "mtime": self.mtime,
            "size": self.size,
            "cloud_only": self.cloud_only,
            "analyzed": self.analyzed,
            "pages": self.pages,
            "first_title": self.first_title,
            "landscape": self.landscape,
            "aspect": self.aspect,
            "chars_per_page": self.chars_per_page,
            "page_size": self.page_size,
            "gdrive_id": self.gdrive_id,
            "own": self.own,
            "is_deck": round(self.is_deck, 2),
            "deck_class": self.deck_class,
            "group_id": self.group_id,
            "talk_id": self.talk_id,
            "match_confidence": None if self.match_confidence is None else round(self.match_confidence, 2),
            "status": self.status,
            "flags": list(self.flags),
        }


# --------------------------------------------------------------------------- walking


_SF_DATALESS = 0x40000000  # macOS: file content lives in iCloud / Drive, not on disk yet


def is_dataless(st: os.stat_result) -> bool:
    """True for iCloud Drive and Google Drive placeholders whose bytes are not local."""
    if getattr(st, "st_flags", 0) & _SF_DATALESS:
        return True
    return st.st_size > 0 and getattr(st, "st_blocks", 1) == 0


def walk_root(root: InputRoot) -> list[FileRecord]:
    """Every deck-like file under the root. Keynote packages (directories) count as files."""
    records: list[FileRecord] = []
    for dirpath, dirnames, filenames in os.walk(root.path, followlinks=False):
        dirnames.sort()
        for name in list(dirnames):
            if name.lower().endswith(".key"):
                dirnames.remove(name)
                records.append(_record(root, Path(dirpath) / name, is_package=True))
        for name in sorted(filenames):
            if name.startswith(".") or name.startswith("._"):
                continue
            if Path(name).suffix.lower() in DECK_EXT:
                records.append(_record(root, Path(dirpath) / name))
    return records


def _record(root: InputRoot, full: Path, is_package: bool = False) -> FileRecord:
    rel = full.relative_to(root.path)
    st = full.stat()
    # .gslides pointers are tiny JSON files that Drive flags as virtual but serves instantly.
    cloud_only = (not is_package) and full.suffix.lower() != ".gslides" and is_dataless(st)
    return FileRecord(
        root=root.name,
        path=str(rel),
        folder=rel.parts[0] if len(rel.parts) > 1 else "",
        type=full.suffix.lower(),
        mtime=date.fromtimestamp(st.st_mtime).isoformat(),
        size=st.st_size if not is_package else sum(p.stat().st_size for p in full.rglob("*") if p.is_file()),
        cloud_only=cloud_only,
    )


# --------------------------------------------------------------------------- analysis


def analyze(rec: FileRecord, full: Path) -> None:
    """Fill content-derived fields. Opens the file, so cloud-only files get downloaded."""
    try:
        if rec.type == ".pdf":
            _analyze_pdf(rec, full)
        elif rec.type == ".pptx":
            _analyze_pptx(rec, full)
        elif rec.type == ".key":
            _analyze_key(rec, full)
        elif rec.type == ".gslides":
            data = json.loads(full.read_text(encoding="utf-8"))
            rec.gdrive_id = data.get("doc_id")
        rec.analyzed = True
    except Exception as exc:  # noqa: BLE001 - one bad file must not stop the walk
        rec.flags.append(f"analyze-error:{type(exc).__name__}")
        rec.analyzed = False


def _analyze_pdf(rec: FileRecord, full: Path) -> None:
    import pymupdf as fitz

    with fitz.open(full) as doc:
        rec.pages = doc.page_count
        if rec.pages == 0:
            return
        sample = [doc[i] for i in range(min(rec.pages, 6))]
        rects = [(round(p.rect.width), round(p.rect.height)) for p in sample]
        w, h = rects[0]
        rec.landscape = w > h
        rec.aspect = round(max(w, h) / max(min(w, h), 1), 3)
        if len(set(rects)) > 1:
            rec.reasons.append("mixed-page-sizes")
        chars = [len(p.get_text("text").strip()) for p in sample]
        rec.chars_per_page = round(sum(chars) / len(chars), 1)
        rec.first_title = _largest_text(sample[0])
        rec.page_size = f"{w}x{h}"


def _largest_text(page: Any) -> str | None:
    spans: list[tuple[float, str]] = []
    for block in page.get_text("dict").get("blocks", []):
        for line in block.get("lines", []):
            for span in line.get("spans", []):
                text = span.get("text", "").strip()
                if text:
                    spans.append((span.get("size", 0.0), text))
    if not spans:
        return None
    biggest = max(size for size, _ in spans)
    title = " ".join(text for size, text in spans if size >= 0.9 * biggest)
    return re.sub(r"\s+", " ", title)[:140] or None


def _analyze_pptx(rec: FileRecord, full: Path) -> None:
    from pptx import Presentation

    prs = Presentation(str(full))
    slides = list(prs.slides)
    rec.pages = len(slides)
    rec.landscape = prs.slide_width > prs.slide_height
    rec.aspect = round(prs.slide_width / max(prs.slide_height, 1), 3)
    if slides:
        title = None
        shapes = slides[0].shapes
        if shapes.title is not None and shapes.title.has_text_frame:
            title = shapes.title.text_frame.text
        if not title:
            for shape in shapes:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    title = shape.text_frame.text
                    break
        if title:
            rec.first_title = re.sub(r"\s+", " ", title).strip()[:140]


def _analyze_key(rec: FileRecord, full: Path) -> None:
    pattern = re.compile(r"^Index/Slide-\d+\.iwa$")
    if full.is_dir():
        names = [str(p.relative_to(full)) for p in full.rglob("*.iwa")]
    else:
        with zipfile.ZipFile(full) as zf:
            names = zf.namelist()
    count = sum(1 for n in names if pattern.match(n))
    rec.pages = count or None
    rec.landscape = True


# --------------------------------------------------------------------------- scoring


def is_own(rec: FileRecord) -> bool:
    """Is this the author's own deck rather than someone else's saved in the trip folder?

    Own: his name in the filename; a Keynote or Google Slides source (others send
    PDFs); a Slides download (-gslide.pptx); or a file inside a talk/presentation
    subfolder. Never own: anything under a templates/downloads/others subfolder.
    """
    parts = [p.lower() for p in Path(rec.path).parts]
    name, dirs = parts[-1], parts[:-1]
    if any(w in d for d in dirs for w in _FOREIGN_DIRS):
        return False
    if _OWN_RE.search(name):
        return True
    if rec.type in (".key", ".gslides") or name.endswith("-gslide.pptx"):
        return True
    return any(w in d for d in dirs[1:] for w in _TALK_DIRS)


def score(rec: FileRecord) -> None:
    """Set is_deck (0-1), deck_class, and reasons from type, name, and content."""
    name = Path(rec.path).name.lower()
    stem = Path(rec.path).stem.lower()
    reasons = list(rec.reasons)
    s = {".key": 0.9, ".gslides": 0.9, ".pptx": 0.85, ".ppt": 0.8, ".pdf": 0.45}[rec.type]

    if any(w in stem for w in _TEMPLATE_WORDS) and not any(w in stem for w in ("talk", "seminar")):
        rec.is_deck, rec.deck_class = 0.2, "template"
        rec.reasons = reasons + ["template-name"]
        return

    hits = [w for w in _POSITIVE_WORDS if w in name]
    if hits:
        s += 0.15
        reasons.append("name+:" + ",".join(hits))
    rec.own = is_own(rec)
    if _OWN_RE.search(name):
        s += 0.1
        reasons.append("own-name")
    neg = [w for w in _NEGATIVE_WORDS if re.search(rf"(?<![a-z]){re.escape(w)}", name)]
    if neg:
        s -= 0.4
        reasons.append("name-:" + ",".join(neg))

    if rec.type == ".pdf" and rec.analyzed and rec.pages:
        if rec.landscape and rec.aspect is not None and any(abs(rec.aspect - a) < 0.06 for a in (1.778, 1.333, 1.6)):
            s += 0.3
            reasons.append("slide-geometry")
        elif not rec.landscape:
            s -= 0.3
            reasons.append("portrait")
        if rec.pages <= 3:
            s -= 0.15
            reasons.append("few-pages")
        elif 8 <= rec.pages <= 150:
            s += 0.1
        if rec.chars_per_page is not None:
            if rec.chars_per_page > 2500:
                s -= 0.3
                reasons.append("dense-text")
            elif rec.chars_per_page < 700:
                s += 0.1
    if rec.type in (".pptx", ".key") and rec.analyzed and rec.pages is not None and rec.pages <= 2:
        s -= 0.2
        reasons.append("few-slides")

    rec.is_deck = max(0.0, min(1.0, s))
    rec.reasons = reasons

    if rec.type == ".pdf" and rec.pages == 1 and rec.page_size:
        w, h = (int(v) for v in rec.page_size.split("x"))
        if min(w, h) > 1200:
            rec.deck_class = "poster"
            rec.reasons.append("poster-geometry")
            return
    if "poster" in name:
        rec.deck_class = "poster"
        return
    if rec.is_deck >= 0.7:
        rec.deck_class = "deck"
    elif rec.is_deck <= 0.3:
        rec.deck_class = "document"
    else:
        rec.deck_class = "unknown"
        rec.flags.append("deck-borderline")


# --------------------------------------------------------------------------- pairing


def split_camel(text: str) -> str:
    """'PrecisionWorkshop-LECM25' -> 'Precision Workshop-LECM25'."""
    return re.sub(r"(?<=[a-z])(?=[A-Z])", " ", text)


def date_key(name: str) -> str | None:
    """'2508.12.Foo' -> '2508.12', '2508.Foo' -> '2508', None without a date prefix."""
    m = _DATE_PREFIX.match(name)
    if not m:
        return None
    return m.group(1) + m.group(2) + (f".{int(m.group(3)):02d}" if m.group(3) else "")


def normalized_stem(rec: FileRecord) -> str:
    stem = split_camel(Path(rec.path).stem).lower()
    stem = _DATE_PREFIX.sub("", stem)
    stem = re.sub(r"\(\d+\)|\bcopy\b|\d+$", " ", stem)
    stem = re.sub(r"[-_.]+", " ", stem)
    stem = _VERSION_TOKENS.sub(" ", stem)
    stem = _NAME_TOKENS.sub(" ", stem)
    return re.sub(r"\s+", " ", stem).strip()


def pair(records: list[FileRecord]) -> dict[str, list[FileRecord]]:
    """Group files that are the same deck in different forms. Returns group_id -> files."""
    candidates = [r for r in records if r.deck_class in ("deck", "unknown", "poster")]
    by_bucket: dict[tuple[str, str], list[FileRecord]] = defaultdict(list)
    for r in candidates:
        by_bucket[(r.root, r.folder)].append(r)

    groups: dict[str, list[FileRecord]] = {}
    counters: dict[str, int] = defaultdict(int)
    for (root, _folder), members in sorted(by_bucket.items()):
        parent = list(range(len(members)))

        def find(i: int) -> int:
            while parent[i] != i:
                parent[i] = parent[parent[i]]
                i = parent[i]
            return i

        stems = [normalized_stem(m) for m in members]
        dates = [date_key(Path(m.path).name) for m in members]
        loose = _folder == ""
        for i in range(len(members)):
            for j in range(i + 1, len(members)):
                if dates[i] and dates[j] and dates[i] != dates[j]:
                    continue  # different date prefixes are different decks
                a, b = stems[i], stems[j]
                if not a or not b:
                    continue
                shorter = min(a, b, key=len)
                same_date = bool(dates[i]) and dates[i] == dates[j]
                contained = (a in b or b in a) and (
                    same_date or (len(shorter.split()) >= 2 and len(shorter) >= 8)
                )
                if loose:
                    same = a == b or fuzz.token_sort_ratio(a, b) >= 90
                else:
                    same = a == b or contained or fuzz.token_sort_ratio(a, b) >= 80
                if same:
                    parent[find(i)] = find(j)
        clusters: dict[int, list[FileRecord]] = defaultdict(list)
        for i, m in enumerate(members):
            clusters[find(i)].append(m)
        for cluster in sorted(clusters.values(), key=lambda c: c[0].path):
            counters[root] += 1
            gid = f"{root}-g{counters[root]:03d}"
            for m in cluster:
                m.group_id = gid
            groups[gid] = cluster
    return groups


# --------------------------------------------------------------------------- matching


def _year_month(text: str) -> str | None:
    m = _DATE_PREFIX.match(text)
    if not m:
        return None
    yy, mm = int(m.group(1)), int(m.group(2))
    if not 1 <= mm <= 12:
        return None
    return f"{2000 + yy:04d}-{mm:02d}"


def group_date(members: list[FileRecord]) -> tuple[str | None, str]:
    """Best guess of the deck's year-month and where it came from."""
    for m in members:
        if m.folder and (ym := _year_month(m.folder)):
            return ym, "folder"
    for m in members:
        if ym := _year_month(Path(m.path).name):
            return ym, "filename"
    dates = sorted(m.mtime for m in members)
    return dates[0][:7], "mtime"


def _months_apart(a: str, b: str) -> int:
    ya, ma = int(a[:4]), int(a[5:7])
    yb, mb = int(b[:4]), int(b[5:7])
    return abs((ya - yb) * 12 + (ma - mb))


def expand_acronyms(text: str) -> str:
    """Lower-case, split on punctuation, and append expansions for known acronyms."""
    words = re.findall(r"[a-z0-9]+", split_camel(text).lower())
    # "tcp22", "fairness19", "lecm25" also count as "tcp", "fairness", "lecm".
    bare = [m.group(1) for w in words if (m := re.fullmatch(r"([a-z]{2,})\d{2,4}", w))]
    words = words + bare
    extra = [_ACRONYMS[w] for w in words if w in _ACRONYMS]
    return " ".join(words + extra)


def group_event_text(members: list[FileRecord]) -> str:
    """Folder and file names only (no slide titles), used to match the CV's event field."""
    parts: set[str] = set()
    for m in members:
        if m.folder:
            parts.add(_DATE_PREFIX.sub("", m.folder))
        parts.add(normalized_stem(m))
    return expand_acronyms(" ".join(p for p in parts if p))


@dataclass
class Candidate:
    talk_id: str
    score: float
    why: str


_W_DATE, _W_EVENT, _W_TITLE = 0.4, 0.35, 0.25


def match_group(
    members: list[FileRecord], talks: list[dict[str, Any]], roots: dict[str, InputRoot]
) -> list[Candidate]:
    ym, ym_source = group_date(members)
    event_text = group_event_text(members)
    titles = [m.first_title for m in members if m.first_title]
    root = roots[members[0].root]
    out: list[Candidate] = []
    for talk in talks:
        tdate = str(talk.get("date", ""))[:7]
        if ym and tdate:
            gap = _months_apart(ym, tdate)
            if gap > 2:
                continue
            date_score = {0: 1.0, 1: 0.6, 2: 0.3}[gap] * (0.6 if ym_source == "mtime" else 1.0)
        else:
            date_score = 0.2
        talk_event = expand_acronyms(str(talk.get("event", "")))
        event_score = fuzz.token_set_ratio(event_text, talk_event) / 100
        title_score = max(
            (fuzz.token_set_ratio(t.lower(), str(talk.get("title", "")).lower()) / 100 for t in titles),
            default=0.0,
        )
        total = _W_DATE * date_score + _W_EVENT * event_score + _W_TITLE * title_score
        if talk.get("era") != root.era:
            total -= 0.05  # files do cross folders (e.g. Slides copies of MIT talks in Drive)
        out.append(Candidate(
            talk["id"], round(total, 3), f"date={date_score:.1f} event={event_score:.2f} title={title_score:.2f}"
        ))
    out.sort(key=lambda c: -c.score)
    return out[:3]


def assign_matches(
    groups: dict[str, list[FileRecord]], talks: list[dict[str, Any]]
) -> dict[str, list[Candidate]]:
    """Set talk_id/match_confidence on unreviewed files. Returns candidates per group."""
    roots = {r.name: r for r in INPUT_ROOTS}
    result: dict[str, list[Candidate]] = {}
    for gid, members in groups.items():
        if not any(m.own and m.deck_class == "deck" for m in members):
            result[gid] = []  # posters, borderline files, and other people's decks are never matched
            continue
        cands = match_group(members, talks, roots)
        result[gid] = cands
        reviewed = [m for m in members if m.status != "unreviewed"]
        if reviewed:
            continue  # the author already decided; keep his talk_id
        best = cands[0] if cands else None
        second = cands[1].score if len(cands) > 1 else 0.0
        confident = best is not None and best.score >= 0.6 and best.score - second >= 0.08
        for m in members:
            m.talk_id = best.talk_id if confident and best else None
            m.match_confidence = best.score if best else None
            if not confident:
                m.flags.append("no-confident-match")
    return result
